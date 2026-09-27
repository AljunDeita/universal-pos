"""
reports_export.py
------------------
Exports tabular report data (sales history, Z-reports, top products) to
PDF, Excel (.xlsx), or Word (.docx) files, so the admin can save/share/print
reports outside the app.
"""

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def export_to_pdf(filepath, title, subtitle, headers, rows, summary_lines=None):
    doc = SimpleDocTemplate(filepath, pagesize=letter,
                             leftMargin=36, rightMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleStyle", parent=styles["Title"], fontSize=16)
    subtitle_style = ParagraphStyle("SubtitleStyle", parent=styles["Normal"], textColor=colors.grey)

    elements = [Paragraph(title, title_style)]
    if subtitle:
        elements.append(Paragraph(subtitle, subtitle_style))
    elements.append(Spacer(1, 14))

    table_data = [headers] + [[str(c) for c in row] for row in rows]
    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e222c")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f5f5")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    elements.append(table)

    if summary_lines:
        elements.append(Spacer(1, 16))
        for label, value in summary_lines:
            elements.append(Paragraph(f"<b>{label}</b> {value}", styles["Normal"]))

    doc.build(elements)
    return filepath


def export_to_excel(filepath, title, subtitle, headers, rows, summary_lines=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(headers)))
    ws.cell(row=1, column=1, value=title).font = Font(size=14, bold=True)

    start_row = 2
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(1, len(headers)))
        ws.cell(row=2, column=1, value=subtitle).font = Font(italic=True, color="666666")
        start_row = 3

    header_row = start_row + 1
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E222C", end_color="1E222C", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for r_idx, row in enumerate(rows, start=header_row + 1):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)

    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 20

    if summary_lines:
        r = header_row + len(rows) + 2
        for label, value in summary_lines:
            ws.cell(row=r, column=1, value=label).font = Font(bold=True)
            ws.cell(row=r, column=2, value=value)
            r += 1

    wb.save(filepath)
    return filepath


def export_to_word(filepath, title, subtitle, headers, rows, summary_lines=None):
    doc = Document()

    h = doc.add_heading(title, level=1)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.runs[0].italic = True

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = str(header)
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                run.bold = True

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)

    if summary_lines:
        doc.add_paragraph("")
        for label, value in summary_lines:
            p = doc.add_paragraph()
            run = p.add_run(f"{label} ")
            run.bold = True
            p.add_run(str(value))

    doc.save(filepath)
    return filepath


def export_report(fmt, filepath, title, subtitle, headers, rows, summary_lines=None):
    """fmt: 'pdf' | 'xlsx' | 'docx'"""
    if fmt == "pdf":
        return export_to_pdf(filepath, title, subtitle, headers, rows, summary_lines)
    elif fmt == "xlsx":
        return export_to_excel(filepath, title, subtitle, headers, rows, summary_lines)
    elif fmt == "docx":
        return export_to_word(filepath, title, subtitle, headers, rows, summary_lines)
    else:
        raise ValueError(f"Unsupported export format: {fmt}")
